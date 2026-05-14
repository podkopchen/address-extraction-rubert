import re
from pathlib import Path

import pandas as pd
import torch
from transformers import AutoTokenizer, AutoModelForTokenClassification

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


# ============================================================
# НАСТРОЙКИ
# ============================================================

BASE_DIR = Path(__file__).resolve().parent

MODEL_DIR = BASE_DIR / "address_ner_rubert_conversational_v1"
INPUT_FILE = BASE_DIR / "Обращения_Адрес_район_округ_очищено.xlsx"
OUTPUT_FILE = BASE_DIR / "Результат_RuBERT_адреса.xlsx"

TEXT_COL = "Обращение заявителя"
MANUAL_ADDRESS_COL = "Выделенный адрес(а)"
RUBERT_ADDRESS_COL = "Адресс rubert"

# Если заголовки в Excel на 3-й строке — header=2.
# Если вдруг не найдет колонки, ниже код сам попробует header=0.
EXCEL_HEADER = 2

MAX_LENGTH = 384
STRIDE = 96


# ============================================================
# ТОКЕНИЗАЦИЯ ДЛЯ PREDICT
# ============================================================

TOKEN_RE = re.compile(
    r"[А-Яа-яЁёA-Za-z]+(?:-[А-Яа-яЁёA-Za-z]+)*|\d+[А-Яа-яЁёA-Za-z]*|№|[^\w\s]"
)


def simple_tokenize(text: str):
    if pd.isna(text):
        return []
    return TOKEN_RE.findall(str(text))


# ============================================================
# BIO / ENTITY POST-PROCESSING
# ============================================================

def repair_bio_sequence(rows):
    """
    Исправляет частый случай:
    B-STREET B-STREET подряд -> B-STREET I-STREET

    Это нужно, например:
    улице B-STREET
    Андропова B-STREET
    превращаем в:
    улице B-STREET
    Андропова I-STREET
    """
    repaired = []
    prev_entity = None

    for idx, word, label in rows:
        if label == "O":
            repaired.append((idx, word, label))
            prev_entity = None
            continue

        if "-" not in label:
            repaired.append((idx, word, label))
            prev_entity = None
            continue

        prefix, entity = label.split("-", 1)

        if prefix == "B" and prev_entity == entity:
            label = f"I-{entity}"

        repaired.append((idx, word, label))
        prev_entity = entity

    return repaired


def extract_entities_from_prediction(rows):
    rows = repair_bio_sequence(rows)

    entities = []
    current = None

    for idx, word, label in rows:
        if label == "O":
            if current:
                entities.append(current)
                current = None
            continue

        if "-" not in label:
            if current:
                entities.append(current)
                current = None
            continue

        prefix, entity_type = label.split("-", 1)

        if prefix == "B":
            if current:
                entities.append(current)

            current = {
                "label": entity_type,
                "tokens": [word],
                "token_ids": [idx],
            }

        elif prefix == "I":
            if current and current["label"] == entity_type:
                current["tokens"].append(word)
                current["token_ids"].append(idx)
            else:
                current = {
                    "label": entity_type,
                    "tokens": [word],
                    "token_ids": [idx],
                }

    if current:
        entities.append(current)

    for ent in entities:
        text = " ".join(ent["tokens"])

        text = text.replace(" .", ".")
        text = text.replace(" ,", ",")
        text = text.replace(" :", ":")
        text = text.replace(" ;", ";")
        text = text.replace(" № ", " №")
        text = text.strip(" ,.;:-")

        ent["text"] = text

    return entities


def clean_component_text(text: str):
    text = str(text)
    text = text.strip(" ,.;:-")
    text = re.sub(r"\s+", " ", text)

    text = text.replace(" .", ".")
    text = text.replace(" ,", ",")
    text = text.replace(" № ", " №")

    return text.strip(" ,.;:-")


def normalize_street(text: str):
    text = clean_component_text(text)

    # улице Андропова -> ул. Андропова
    text = re.sub(r"^(улица|улице|улицу|ул\.?|ул)\s+", "ул. ", text, flags=re.IGNORECASE)

    # проспект / пр-т
    text = re.sub(r"^(проспект|пр-т|пр\.?)\s+", "пр. ", text, flags=re.IGNORECASE)

    # проезд / пр-д
    text = re.sub(r"^(проезд|пр-д)\s+", "пр-д ", text, flags=re.IGNORECASE)

    # переулок
    text = re.sub(r"^(переулок|пер\.?)\s+", "пер. ", text, flags=re.IGNORECASE)

    # шоссе
    text = re.sub(r"^(шоссе)\s+", "ш. ", text, flags=re.IGNORECASE)

    # бульвар
    text = re.sub(r"^(бульвар|б-р)\s+", "б-р ", text, flags=re.IGNORECASE)

    return text


def normalize_house(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(дом|д\.?|д|№)\s+", "", text, flags=re.IGNORECASE)
    return f"д. {text}" if text else ""


def normalize_corpus(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(корпус|корп\.?|к\.?|к)\s+", "", text, flags=re.IGNORECASE)
    return f"к. {text}" if text else ""


def normalize_building(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(строение|стр\.?|с\.?|с)\s+", "", text, flags=re.IGNORECASE)
    return f"с. {text}" if text else ""


def normalize_entrance(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(подъезд|под\.?|п\.?|п)\s+", "", text, flags=re.IGNORECASE)
    return f"подъезд {text}" if text else ""


def normalize_flat(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(квартира|кв\.?|кв)\s+", "", text, flags=re.IGNORECASE)
    return f"кв. {text}" if text else ""


def normalize_city(text: str):
    text = clean_component_text(text)
    text = re.sub(r"^(город|г\.?|г)\s+", "", text, flags=re.IGNORECASE)
    return f"г. {text}" if text else ""


def entities_to_address(entities):
    """
    Собирает нормализованную строку адреса из компонент.

    Важно:
    если CITY не найден, г. Москва НЕ добавляем.
    """
    by_label = {}

    for ent in entities:
        label = ent["label"]
        value = ent["text"]

        if not value:
            continue

        by_label.setdefault(label, [])
        by_label[label].append(value)

    # Если модель нашла несколько адресов, пока собираем все компоненты подряд.
    # Для первичной проверки таблицы этого достаточно.
    parts = []

    for region in by_label.get("REGION", []):
        parts.append(clean_component_text(region))

    for city in by_label.get("CITY", []):
        parts.append(normalize_city(city))

    for settlement in by_label.get("SETTLEMENT", []):
        parts.append(clean_component_text(settlement))

    for district in by_label.get("DISTRICT", []):
        parts.append(clean_component_text(district))

    for street in by_label.get("STREET", []):
        parts.append(normalize_street(street))

    for house in by_label.get("HOUSE", []):
        parts.append(normalize_house(house))

    for corpus in by_label.get("CORPUS", []):
        parts.append(normalize_corpus(corpus))

    for building in by_label.get("BUILDING", []):
        parts.append(normalize_building(building))

    for entrance in by_label.get("ENTRANCE", []):
        parts.append(normalize_entrance(entrance))

    for flat in by_label.get("FLAT", []):
        parts.append(normalize_flat(flat))

    parts = [p for p in parts if p]

    if not parts:
        return ""

    return ", ".join(parts)


# ============================================================
# PREDICT
# ============================================================

def predict_rows_for_text(text, tokenizer, model, id2label, device):
    words = simple_tokenize(text)

    if not words:
        return []

    # Для длинных текстов используем overflow + stride.
    # padding=True обязательно, иначе разные чанки 384/338 не соберутся в tensor.
    enc = tokenizer(
        words,
        is_split_into_words=True,
        return_tensors="pt",
        truncation=True,
        max_length=MAX_LENGTH,
        stride=STRIDE,
        return_overflowing_tokens=True,
        padding=True,
    )

    all_rows = []
    seen_word_ids = set()

    input_ids = enc["input_ids"]
    attention_mask = enc["attention_mask"]

    for chunk_idx in range(input_ids.shape[0]):
        chunk_inputs = {
            "input_ids": input_ids[chunk_idx:chunk_idx + 1].to(device),
            "attention_mask": attention_mask[chunk_idx:chunk_idx + 1].to(device),
        }

        if "token_type_ids" in enc:
            chunk_inputs["token_type_ids"] = enc["token_type_ids"][chunk_idx:chunk_idx + 1].to(device)

        word_ids = enc.word_ids(batch_index=chunk_idx)

        with torch.no_grad():
            out = model(**chunk_inputs)

        pred_ids = out.logits.argmax(dim=-1)[0].detach().cpu().numpy().tolist()

        for token_idx, word_idx in enumerate(word_ids):
            if word_idx is None:
                continue

            # Берем только первый subword каждого исходного токена.
            # Если токен уже встретился в предыдущем overlapping chunk — пропускаем.
            if word_idx in seen_word_ids:
                continue

            seen_word_ids.add(word_idx)

            word = words[word_idx]
            label = id2label[int(pred_ids[token_idx])]
            all_rows.append((word_idx + 1, word, label))

    all_rows.sort(key=lambda x: x[0])
    return all_rows


def predict_address(text, tokenizer, model, id2label, device):
    rows = predict_rows_for_text(text, tokenizer, model, id2label, device)
    entities = extract_entities_from_prediction(rows)
    address = entities_to_address(entities)

    raw_entities = "; ".join(
        f"{ent['label']}={ent['text']}" for ent in entities
    )

    return address, raw_entities


# ============================================================
# СРАВНЕНИЕ С РУЧНЫМ АДРЕСОМ
# ============================================================

def normalize_for_compare(text: str):
    if pd.isna(text):
        return ""

    text = str(text).lower().strip()

    if text in {"", "отсутствует", "нет", "nan", "none"}:
        return ""

    replacements = {
        "ё": "е",
        "город ": "г ",
        "г. ": "г ",
        "улица ": "ул ",
        "улице ": "ул ",
        "ул. ": "ул ",
        "дом ": "д ",
        "д. ": "д ",
        "корпус ": "к ",
        "корп. ": "к ",
        "к. ": "к ",
        "строение ": "с ",
        "стр. ": "с ",
        "квартира ": "кв ",
        "кв. ": "кв ",
    }

    for src, dst in replacements.items():
        text = text.replace(src, dst)

    text = re.sub(r"[^\wа-яА-ЯёЁ0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()

    return text


def compare_addresses(manual, rubert):
    manual_norm = normalize_for_compare(manual)
    rubert_norm = normalize_for_compare(rubert)

    if not manual_norm and not rubert_norm:
        return "оба пусто"

    if manual_norm and not rubert_norm:
        return "rubert пусто"

    if not manual_norm and rubert_norm:
        return "rubert нашел, вручную пусто"

    if manual_norm == rubert_norm:
        return "полное совпадение"

    if manual_norm in rubert_norm:
        return "rubert шире"

    if rubert_norm in manual_norm:
        return "ручной шире"

    manual_tokens = set(manual_norm.split())
    rubert_tokens = set(rubert_norm.split())

    if not manual_tokens or not rubert_tokens:
        return "разное"

    intersection = manual_tokens & rubert_tokens
    union = manual_tokens | rubert_tokens
    jaccard = len(intersection) / len(union)

    if jaccard >= 0.75:
        return "похоже"
    elif jaccard >= 0.45:
        return "частично похоже"
    else:
        return "разное"


# ============================================================
# EXCEL FORMAT
# ============================================================

def format_excel(path):
    wb = load_workbook(path)
    ws = wb.active
    ws.title = "RuBERT адреса"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9"),
    )

    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(
                horizontal="left",
                vertical="center",
                wrap_text=True,
            )
            cell.border = thin_border

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    ws.auto_filter.ref = ws.dimensions
    ws.freeze_panes = "A2"

    widths = {
        "A": 12,
        "B": 85,
        "C": 45,
        "D": 45,
        "E": 28,
        "F": 70,
    }

    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = widths.get(col_letter, 30)

    ws.row_dimensions[1].height = 35

    for row_num in range(2, ws.max_row + 1):
        ws.row_dimensions[row_num].height = 85

    wb.save(path)


# ============================================================
# MAIN
# ============================================================

def read_excel_smart(path):
    df_try = pd.read_excel(path, header=EXCEL_HEADER)

    if TEXT_COL in df_try.columns and MANUAL_ADDRESS_COL in df_try.columns:
        return df_try

    df_try = pd.read_excel(path, header=0)

    if TEXT_COL in df_try.columns and MANUAL_ADDRESS_COL in df_try.columns:
        return df_try

    raise ValueError(
        "Не нашел нужные колонки. "
        f"Нужны: {TEXT_COL!r}, {MANUAL_ADDRESS_COL!r}. "
        f"Фактические колонки: {list(df_try.columns)}"
    )


def main():
    print("BASE_DIR:", BASE_DIR)
    print("MODEL_DIR:", MODEL_DIR)
    print("INPUT_FILE:", INPUT_FILE)
    print("OUTPUT_FILE:", OUTPUT_FILE)

    if not MODEL_DIR.exists():
        raise FileNotFoundError(f"Не найдена папка модели: {MODEL_DIR}")

    if not INPUT_FILE.exists():
        raise FileNotFoundError(f"Не найден Excel-файл: {INPUT_FILE}")

    print()
    print("Загружаю модель...")

    tokenizer = AutoTokenizer.from_pretrained(MODEL_DIR)
    model = AutoModelForTokenClassification.from_pretrained(MODEL_DIR)

    device = torch.device("cuda" if torch.cuda.is_available() else "cpu")
    model.to(device)
    model.eval()

    id2label = model.config.id2label

    # Иногда ключи могут быть строками, приводим к int.
    id2label = {int(k): v for k, v in id2label.items()}

    print("Device:", device)
    print("Labels:", id2label)

    print()
    print("Читаю Excel...")

    df = read_excel_smart(INPUT_FILE)

    print("Строк:", len(df))
    print("Колонки:", list(df.columns))

    output_rows = []

    for i, row in df.iterrows():
        text = row.get(TEXT_COL, "")
        manual_address = row.get(MANUAL_ADDRESS_COL, "")

        rubert_address, raw_entities = predict_address(
            text=text,
            tokenizer=tokenizer,
            model=model,
            id2label=id2label,
            device=device,
        )

        compare_status = compare_addresses(manual_address, rubert_address)

        output_rows.append({
            "№": i + 1,
            "Обращение заявителя": text,
            "Выделенный адрес(а)": manual_address,
            "Адресс rubert": rubert_address,
            "Сравнение": compare_status,
            "Компоненты rubert": raw_entities,
        })

        if (i + 1) % 50 == 0:
            print(f"Обработано: {i + 1}/{len(df)}")

    result_df = pd.DataFrame(output_rows)

    result_df.to_excel(OUTPUT_FILE, index=False)
    format_excel(OUTPUT_FILE)

    print()
    print("Готово:", OUTPUT_FILE)
    print()
    print("Статистика сравнения:")
    print(result_df["Сравнение"].value_counts(dropna=False))


if __name__ == "__main__":
    main()
